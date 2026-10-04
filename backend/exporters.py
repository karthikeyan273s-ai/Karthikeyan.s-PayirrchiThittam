from io import BytesIO
import base64
import re
import html

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from fpdf import FPDF

from backend.utils import (
    sanitize_text,
    terms_list
)


def _logo_bytes(value):

    if not value:
        return None

    try:

        return base64.b64decode(
            value.split(",", 1)[-1],
            validate=True
        )

    except Exception:

        return None


def format_txt(text):

    return sanitize_text(text).encode(
        "utf-8"
    )


def _shade(cell, fill="E9EEF5"):

    tc_pr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")

    shd.set(
        qn("w:fill"),
        fill
    )

    tc_pr.append(shd)


def format_docx(
    text,
    doc_type,
    terms="",
    logo_base64=None
):

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    doc.styles["Normal"].font.name = (
        "Times New Roman"
    )

    doc.styles["Normal"].font.size = Pt(11)

    logo = _logo_bytes(logo_base64)

    if logo:

        paragraph = doc.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.add_run().add_picture(
            BytesIO(logo),
            width=Inches(1.25)
        )

    paragraph = doc.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run(
        sanitize_text(
            doc_type
        ).upper()
    )

    run.bold = True
    run.font.size = Pt(16)

    doc.add_paragraph()

    for raw in sanitize_text(text).split("\n"):

        line = raw.strip()

        if not line:

            doc.add_paragraph()
            continue

        paragraph = doc.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(6)

        heading = (
            len(line) <= 90
            and (
                line.isupper()
                or line.lower().startswith(
                    (
                        "section ",
                        "article ",
                        "recitals",
                        "whereas"
                    )
                )
                or bool(
                    re.match(
                        r"^\d+\.",
                        line
                    )
                )
            )
        )

        run = paragraph.add_run(line)

        run.bold = heading

        if heading:
            run.font.size = Pt(12)

    parsed_terms = terms_list(terms)

    if parsed_terms:

        doc.add_page_break()

        heading = doc.add_paragraph()

        run = heading.add_run(
            "KEY TERMS"
        )

        run.bold = True
        run.font.size = Pt(13)

        table = doc.add_table(
            rows=1,
            cols=2
        )

        table.style = "Table Grid"

        table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = (
            "Term / Condition"
        )

        for cell in table.rows[0].cells:

            _shade(cell)

            for run in (
                cell.paragraphs[0].runs
            ):

                run.bold = True

        for i, term in enumerate(
            parsed_terms,
            1
        ):

            cells = table.add_row().cells

            cells[0].text = str(i)
            cells[1].text = term

    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer.add_run(
        "LegalEase - AI-assisted draft | "
        "Review with a qualified legal professional "
        "before use."
    )

    output = BytesIO()

    doc.save(output)

    return output.getvalue()


class LegalPDF(FPDF):

    def __init__(self, logo=None):

        super().__init__()

        self.logo = logo

    def header(self):

        if self.logo:

            try:

                self.image(
                    BytesIO(self.logo),
                    x=92,
                    y=8,
                    w=26
                )

                self.set_y(36)

            except Exception:

                self.set_y(12)

        else:

            self.set_y(12)

        self.set_font(
            "Helvetica",
            "B",
            9
        )

        self.cell(
            0,
            6,
            "LegalEase",
            align="C"
        )

        self.ln(7)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            "",
            7
        )

        self.cell(
            0,
            8,
            "LegalEase - AI-assisted draft. "
            "Not a substitute for legal advice.",
            align="C"
        )


def format_pdf(
    text,
    doc_type,
    terms="",
    logo_base64=None
):

    pdf = LegalPDF(
        _logo_bytes(logo_base64)
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        "B",
        16
    )

    pdf.multi_cell(
        0,
        9,
        sanitize_text(
            doc_type
        ).upper(),
        align="C"
    )

    pdf.ln(4)

    pdf.set_font(
        "Helvetica",
        "",
        11
    )

    for raw in sanitize_text(text).split("\n"):

        line = raw.strip()

        if not line:

            pdf.ln(4)
            continue

        heading = (
            len(line) <= 90
            and (
                line.isupper()
                or line.lower().startswith(
                    (
                        "section ",
                        "article ",
                        "recitals",
                        "whereas"
                    )
                )
                or bool(
                    re.match(
                        r"^\d+\.",
                        line
                    )
                )
            )
        )

        pdf.set_font(
            "Helvetica",
            "B" if heading else "",
            12 if heading else 11
        )

        pdf.multi_cell(
            0,
            7 if heading else 6,
            line
        )

    parsed_terms = terms_list(terms)

    if parsed_terms:

        pdf.add_page()

        pdf.set_font(
            "Helvetica",
            "B",
            13
        )

        pdf.cell(
            0,
            8,
            "KEY TERMS"
        )

        pdf.ln(10)

        pdf.set_font(
            "Helvetica",
            "",
            10
        )

        for i, term in enumerate(
            parsed_terms,
            1
        ):

            pdf.multi_cell(
                0,
                6,
                f"{i}. {term}"
            )

            pdf.ln(1)

    return bytes(
        pdf.output(
            dest="S"
        )
    )


def format_html_preview(text):

    blocks = []

    for raw in sanitize_text(text).split(
        "\n"
    ):

        line = raw.strip()

        if not line:
            continue

        safe = html.escape(line)

        heading = (
            len(line) <= 90
            and (
                line.isupper()
                or line.lower().startswith(
                    (
                        "section ",
                        "article ",
                        "recitals",
                        "whereas"
                    )
                )
                or bool(
                    re.match(
                        r"^\d+\.",
                        line
                    )
                )
            )
        )

        if heading:

            blocks.append(
                f"<h3>{safe}</h3>"
            )

        else:

            blocks.append(
                f"<p>{safe}</p>"
            )

    return "".join(blocks)